"""`services/xlsx` — in-place edit round-trip on a temp workbook."""

from openpyxl import Workbook

from frontend.services import xlsx


def _book(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    ws.append(["name", "count", "ratio"])
    ws.append(["a", 3, 0.5])
    ws.append(["b", 7, 1.25])
    p = tmp_path / "spec.xlsx"
    wb.save(p)
    return p


def test_grid_and_edit_roundtrip(tmp_path):
    p = _book(tmp_path)
    wb = xlsx.open_book(p)
    assert xlsx.sheet_names(wb) == ["S1"]

    g = xlsx.grid(wb, "S1")
    assert g[0] == ["name", "count", "ratio"]
    assert g[1] == ["a", "3", "0.5"]

    # int cell keeps int type; string stays string; float stays float
    xlsx.set_cell(wb, "S1", 1, 1, "42")      # was int 3
    xlsx.set_cell(wb, "S1", 1, 0, "aa")      # was str
    xlsx.set_cell(wb, "S1", 2, 2, "9,9")     # comma -> must stay string
    xlsx.save_book(wb, p)

    wb2 = xlsx.open_book(p)
    ws = wb2["S1"]
    assert ws.cell(2, 2).value == 42 and isinstance(ws.cell(2, 2).value, int)
    assert ws.cell(2, 1).value == "aa"
    assert ws.cell(3, 3).value == "9,9"


def test_add_row_col(tmp_path):
    wb = xlsx.open_book(_book(tmp_path))
    r0 = wb["S1"].max_row
    c0 = wb["S1"].max_column
    xlsx.add_row(wb, "S1")
    xlsx.add_col(wb, "S1")
    assert wb["S1"].max_row == r0 + 1
    assert wb["S1"].max_column == c0 + 1
