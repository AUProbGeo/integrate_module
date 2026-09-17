"""Small in-place spreadsheet editing on top of openpyxl.

The live ``Workbook`` is kept in the session (single-user desktop tool), so
edits touch only the cells the user changes — other cells, sheets, formulas
and styles are written back untouched on save. Intended for the small
geoprior1d Excel specs (a few sheets of ~10×10), not big data grids.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

MAX_ROWS = 200
MAX_COLS = 40


def open_book(path: str | Path) -> Workbook:
    return load_workbook(path, data_only=False)


def sheet_names(wb: Workbook) -> list[str]:
    return list(wb.sheetnames)


def grid(wb: Workbook, sheet: str) -> list[list[str]]:
    """Rectangular list-of-rows of display strings for *sheet* (headerless)."""
    ws = wb[sheet]
    n_rows = min(ws.max_row or 1, MAX_ROWS)
    n_cols = min(ws.max_column or 1, MAX_COLS)
    out = []
    for r in range(1, n_rows + 1):
        row = []
        for c in range(1, n_cols + 1):
            v = ws.cell(row=r, column=c).value
            row.append("" if v is None else str(v))
        out.append(row)
    return out


def col_labels(n: int) -> list[str]:
    return [get_column_letter(i + 1) for i in range(n)]


def _coerce(text: str, old):
    """Cast *text* to match *old*'s type where sensible; keep strings as-is."""
    t = text.strip()
    if t == "":
        return None
    if isinstance(old, bool):
        return t.lower() in ("1", "true", "yes")
    if isinstance(old, int) and not isinstance(old, bool):
        try:
            return int(t)
        except ValueError:
            pass
    if isinstance(old, float):
        try:
            return float(t)
        except ValueError:
            pass
    if old is None:  # brand-new cell: best-effort number, else text
        for cast in (int, float):
            try:
                return cast(t)
            except ValueError:
                pass
    return t


def set_cell(wb: Workbook, sheet: str, r: int, c: int, text: str) -> None:
    """Set 0-indexed cell (r, c) of *sheet*, coercing against its old value."""
    if not (0 <= r < MAX_ROWS and 0 <= c < MAX_COLS):
        raise ValueError("cell out of bounds")
    ws = wb[sheet]
    cell = ws.cell(row=r + 1, column=c + 1)
    cell.value = _coerce(text, cell.value)


def add_row(wb: Workbook, sheet: str) -> None:
    ws = wb[sheet]
    if (ws.max_row or 0) >= MAX_ROWS:
        return
    ws.append([None] * (ws.max_column or 1))


def add_col(wb: Workbook, sheet: str) -> None:
    ws = wb[sheet]
    c = (ws.max_column or 0) + 1
    if c > MAX_COLS:
        return
    ws.cell(row=1, column=c).value = None


def save_book(wb: Workbook, path: str | Path) -> None:
    wb.save(path)


def new_book(sheet: str = "Sheet1") -> Workbook:
    wb = Workbook()
    wb.active.title = sheet
    return wb
